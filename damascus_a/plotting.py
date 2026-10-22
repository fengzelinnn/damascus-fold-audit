"""A1-A5 empty layouts and explicitly NON-PAPER fixture previews.

Exports vector PDF/SVG and native editable PowerPoint charts, axes and tables.
No benchmark is run here. Publication export needs both approved input data
and the requested Libertine/newtx typography; missing fonts fail closed.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
import json
import os
from pathlib import Path
import shutil
import subprocess
import textwrap

from .aggregate import STAGES, ValidationError, aggregate, load_records

EXPERIMENTS = ("A1", "A2", "A3", "A4", "A5")
TITLES = {
    "A1": "A1  Serialized communication versus file size",
    "A2": "A2  Entire Verify phase: cold and cached",
    "A3": "A3  Average cost across successive NEW epochs",
    "A4": "A4  Two local platforms: resources and tested failures",
    "A5": "A5  Baseline semantics and provenance",
}
GRAY = ("#161616", "#4a4a4a", "#7c7c7c", "#a5a5a5", "#c4c4c4", "#dedede", "#eeeeee")
DASHES = ("-", "--", "-.", ":")
HATCHES = ("", "///", "\\\\", "xx", "..", "++", "oo")
STAGE_LABELS = ("Parse / context", "Challenges", "Generator replay", "Commitment / parity", "Norm", "Opening", "Other (explicit)")
MIB = 2 ** 20


@dataclass(frozen=True)
class Typography:
    publication_ready: bool
    family: str
    missing: tuple[str, ...]


def check_typography() -> Typography:
    """Require exact desktop text font and TeX packages, without silent fallback."""
    # Avoid matplotlib attempting to create a read-only home cache.
    os.environ.setdefault("MPLCONFIGDIR", "/tmp/damascus-a-matplotlib")
    from matplotlib import font_manager
    missing = []
    try:
        path = font_manager.findfont(font_manager.FontProperties(family="Linux Libertine O"), fallback_to_default=False)
        available = Path(path).is_file()
    except ValueError:
        available = False
    if not available:
        missing.append("Linux Libertine O desktop font")
    for command in ("latex", "kpsewhich"):
        if shutil.which(command) is None:
            missing.append(command)
    if shutil.which("kpsewhich"):
        for package in ("libertine.sty", "newtxmath.sty"):
            checked = subprocess.run(["kpsewhich", package], capture_output=True, text=True, timeout=10, check=False)
            if checked.returncode != 0 or not checked.stdout.strip():
                missing.append(package)
    return Typography(not missing, "Linux Libertine O" if not missing else "DejaVu Serif", tuple(missing))


def _status(aggregated: list[dict], publication: bool, typography: Typography) -> tuple[str, str]:
    kinds = set()
    for figure in aggregated:
        metadata = figure["metadata"]
        kinds.update(metadata.get("data_kinds", []))
        if metadata.get("data_kind"):
            kinds.add(metadata["data_kind"])
    if "fixture" in kinds:
        banner = "NON-PAPER — SYNTHETIC TEST FIXTURE — NOT MEASUREMENTS"
    elif all(a["empty"] for a in aggregated):
        banner = "EMPTY LAYOUT — NO MEASUREMENTS / NO PAPER RESULTS"
    elif publication:
        banner = "DATA PRECONDITIONS CHECKED — SEE PROVENANCE AND EXTERNAL PROTOCOL REVIEW"
    else:
        banner = "PREVIEW — INPUT LOGS ONLY — NOT A PUBLICATION APPROVAL"
    font = "Libertine + newtx math" if typography.publication_ready else "PREVIEW FONT FALLBACK: DejaVu Serif / DejaVu math; Libertine + newtx unavailable"
    return banner, font


def _panels(figure: dict) -> list[dict]:
    """Render-neutral numeric series used by both vector and native PPT charts."""
    name, rows = figure["experiment"], figure["rows"]
    if name == "A1":
        series = []
        if rows:
            xs = [r["file_bytes"] / MIB for r in rows]
            series = [{"label": "Full proof (includes opening)", "x": xs, "y": [r["full_proof_bytes"] for r in rows]}, {"label": "Registration", "x": xs, "y": [r["registration_bytes"] for r in rows]}]
            for index in range(max(len(r["round_bytes"]) for r in rows)):
                pairs = [(r["file_bytes"] / MIB, r["round_bytes"][index]) for r in rows if index < len(r["round_bytes"])]
                series.append({"label": f"Round record {index + 1}", "x": [p[0] for p in pairs], "y": [p[1] for p in pairs]})
        return [{"title": "Actual artifact bytes", "x_label": "File size (MiB)", "y_label": "Serialized bytes", "series": series}]
    if name == "A2":
        return [{"title": "Cold / cached pair at each tested size", "x_label": "File size (MiB) / cache state", "y_label": "Entire Verify (ms)", "categories": [f"{r['file_bytes'] / MIB:g}\n{'cold' if r['cache'] == 'cold' else 'cached'}" for r in rows], "series": [{"label": label, "y": [r["stages_ns"][stage] / 1e6 for r in rows]} for stage, label in zip(STAGES, STAGE_LABELS)] if rows else []}]
    if name == "A3":
        return [{"title": "One registered file; a new certified context and proof each audit", "x_label": "Successive NEW epoch audits", "y_label": "Mean cost per audit (ms)", "series": [{"label": "(Setup + sum(Prove + Verify)) / audits", "x": [r["audit_count"] for r in rows], "y": [r["with_setup_ns"] / 1e6 for r in rows]}, {"label": "sum(Prove + Verify) / audits (setup excluded)", "x": [r["audit_count"] for r in rows], "y": [r["without_setup_ns"] / 1e6 for r in rows]}] if rows else []}]
    if name == "A4":
        platforms = sorted({r["hardware_id"] for r in rows}) or ["Local platform 1 — identity pending", "Local platform 2 — identity pending"]
        panels = []
        for platform in platforms:
            selected = [r for r in rows if r["hardware_id"] == platform]
            series = []
            for metric, label in (("peak_rss_bytes", "Peak process RSS"), ("sampled_peak_vram_bytes", "Sampled VRAM maximum")):
                # Preserve each repetition rather than silently averaging peaks.
                for run_id, repetition in sorted({(r["run_id"], r["repetition"]) for r in selected}):
                    values = [r for r in selected if (r["run_id"], r["repetition"]) == (run_id, repetition) and r["status"] == "ok" and r["resources"].get(metric) is not None]
                    if values:
                        series.append({"label": f"{label} ({run_id}, rep {repetition})", "x": [r["file_bytes"] / MIB for r in values], "y": [r["resources"][metric] / MIB for r in values]})
            failures = [{"x": r["file_bytes"] / MIB, "status": r["status"], "reason": r["failure_reason"], "run_id": r["run_id"], "repetition": r["repetition"]} for r in selected if r["status"] != "ok"]
            panels.append({"title": platform, "x_label": "File size (MiB)", "y_label": "Memory (MiB)", "series": series, "failures": failures, "vram_missing": bool(selected) and all(r["resources"].get("sampled_peak_vram_bytes") is None for r in selected)})
        return panels
    return []


def _comparison_rows(figure: dict) -> list[list[str]]:
    if not figure["rows"]:
        return [
            ["ADJ21\nPublic low-storage PoR", "Point-in-time retrievability; no duration guarantee", "LIP / discrete log / collision resistance; local profile pending", "USENIX Security 2021; see docs/baselines.md. No imported timing results", "No local measurements"],
            ["ePoSt\nStorage-time semantics", "Public, stateless storage-time (abstract); exact theorem unverified", "Primary full text unavailable; concrete assumptions / parameters unverified", "TIFS 18 (2023); abstract evidence only. See docs/baselines.md", "No local measurements; published hardware unverified"],
            ["Damascus\nCurrent manuscript", "No observation", "D1–D6 unresolved", "No paper results", "Two local platforms pending"],
        ]
    rendered = []
    for row in figure["rows"]:
        comp, prov = row["comparison"], row["provenance"]
        rendered.append([
            f"{row['scheme']} {row['scheme_version']}\n{comp['protocol_family']}",
            comp["guarantee"],
            f"{comp['security_model']}\n{comp['parameter_notes']}\nProfile: {row['profile']}",
            f"{row['data_kind']}: {comp['source']}\n{comp['cost_scope']}\n{prov['source_id']}",
            f"{row['hardware_id']}\n{prov['hardware_description']}\nBackend: {row['backend']['actual']}",
        ])
    return rendered


def _notes(name: str) -> str:
    return {
        "A1": "Full proof, each round record and registration are separate series. Counts include actual framing and opening bytes; no formula-based estimates.",
        "A2": "Stack covers all elapsed Verify time, including opening and explicit residual Other. Cold/cached refer to declared cache policy, with hardware/profile held fixed.",
        "A3": "Setup is charged once per independent sequence. Each point averages cumulative Prove + entire Verify over successive NEW epochs; cached replay of one proof is rejected.",
        "A4": "Failure markers identify tested file sizes only. RSS is process peak; VRAM is an observed sampling maximum. Missing VRAM stays unavailable, never zero-filled.",
        "A5": "Public low-storage PoR and storage-time guarantees are distinct. Sources, assumptions, profiles and hardware stay separate; seconds from different systems are not ranked.",
    }[name]


def _identity(figure: dict) -> str:
    meta = figure["metadata"]
    if figure["empty"]:
        return "Protocol profile, parameters and local platform identity remain pending. A5 source review: docs/baselines.md."
    if figure["experiment"] == "A5":
        return "Evidence kind, semantic scope, source and hardware are retained per row. No pooled timing comparison."
    params = f"schedule {meta['parameter_schedule_id']} (size mapping in manifest)" if "parameter_schedule_id" in meta else meta["parameter_id"]
    return f"Profile: {meta['profile']} | Parameters: {params} | Hardware: {meta.get('hardware_id', 'two separate panels')} | Backend: {meta['backend']}"


def _tex_literal(text: str) -> str:
    """Escape literal evidence labels once; never interpret source text as TeX."""
    replacements = {
        "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$",
        "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}",
        "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
        "|": r"\textbar{}", "<": r"\textless{}", ">": r"\textgreater{}",
    }
    return "".join(replacements.get(char, char) for char in str(text))


def _vector(figure: dict, output: Path, banner: str, font_note: str, typography: Typography, publication: bool) -> None:
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    rc = {"font.family": typography.family, "font.size": 9, "axes.spines.top": False, "axes.spines.right": False, "svg.fonttype": "none", "pdf.fonttype": 42, "axes.grid": True, "grid.alpha": 0.2, "mathtext.fontset": "dejavuserif"}
    if publication:
        rc.update({"text.usetex": True, "text.latex.preamble": r"\usepackage{libertine}\usepackage[libertine]{newtxmath}"})
    literal = _tex_literal if publication else str
    with plt.rc_context(rc):
        name = figure["experiment"]
        fig = plt.figure(figsize=(12, 6.75), facecolor="white")
        fig.text(0.04, 0.95, literal(TITLES[name]), fontsize=17, weight="bold", va="top")
        fig.text(0.04, 0.895, literal(banner), fontsize=9, weight="bold", color="#333333")
        fig.text(0.04, 0.855, literal(_identity(figure)), fontsize=7, color="#4a4a4a")
        fig.text(0.04, 0.035, literal(font_note), fontsize=7, color="#4a4a4a")
        fig.text(0.04, 0.075, literal(textwrap.fill(_notes(name), 148)), fontsize=8, va="bottom")
        panels = _panels(figure)
        if name == "A5":
            ax = fig.add_axes((0.04, 0.22, 0.92, 0.59))
            ax.axis("off")
            data = [[literal(textwrap.fill(cell, 28)) for cell in row] for row in _comparison_rows(figure)]
            table = ax.table(cellText=data, colLabels=[literal(label) for label in ("Scheme / family", "Semantics / guarantee", "Security / parameters", "Evidence / source / scope", "Hardware / backend")], loc="center", cellLoc="left", bbox=(0, 0, 1, 1))
            table.auto_set_font_size(False)
            table.set_fontsize(8 if len(data) <= 3 else 6)
            for (row, col), cell in table.get_celld().items():
                cell.set_edgecolor("#777777")
                cell.set_facecolor("#dedede" if row == 0 else ("#f4f4f4" if row % 2 else "#ffffff"))
                if row == 0:
                    cell.set_text_props(weight="bold")
        else:
            for index, panel in enumerate(panels):
                width = 0.88 if len(panels) == 1 else 0.39
                left = 0.08 if len(panels) == 1 else 0.075 + index * 0.49
                ax = fig.add_axes((left, 0.29, width, 0.49))
                ax.set_title(literal(panel["title"]), fontsize=10, pad=12)
                ax.set_xlabel(literal(panel["x_label"]))
                ax.set_ylabel(literal(panel["y_label"]))
                series = panel["series"]
                if not series:
                    ax.set_xticks([])
                    ax.set_yticks([])
                    ax.text(0.5, 0.5, literal("NO DATA\nAwaiting authorized local measurements"), ha="center", va="center", transform=ax.transAxes, color="#666666", fontsize=11)
                elif name == "A2":
                    positions = list(range(len(panel["categories"])))
                    bottom = [0.0] * len(positions)
                    for i, s in enumerate(series):
                        ax.bar(positions, s["y"], bottom=bottom, label=literal(s["label"]), color=GRAY[i % len(GRAY)], edgecolor="black", linewidth=0.5, hatch=HATCHES[i % len(HATCHES)], width=0.7)
                        bottom = [a + b for a, b in zip(bottom, s["y"])]
                    ax.set_xticks(positions, [literal(label) for label in panel["categories"]])
                    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.24), ncols=4, fontsize=7, frameon=False)
                else:
                    for i, s in enumerate(series):
                        ax.plot(s["x"], s["y"], label=literal(s["label"]), color=GRAY[i % 4], linestyle=DASHES[i % len(DASHES)], marker=("o", "s", "^", "D")[i % 4], markersize=4, linewidth=1.2, markerfacecolor="white")
                    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.2), ncols=1 if name == "A4" else 2, fontsize=7, frameon=False)
                for i, failure in enumerate(panel.get("failures", [])):
                    ax.plot([failure["x"]], [0.02], transform=ax.get_xaxis_transform(), color="black", marker="x", clip_on=False)
                    ax.annotate(literal(f"{failure['status']} @ {failure['x']:g} MiB"), xy=(failure["x"], 0.04 + (i % 3) * 0.07), xycoords=("data", "axes fraction"), fontsize=7, ha="center")
                if panel.get("vram_missing"):
                    ax.text(0.98, 0.96, literal("VRAM unavailable"), transform=ax.transAxes, ha="right", va="top", fontsize=7)
                ax.set_ylim(bottom=0)
        # Do not rasterize any artist; both outputs remain vector graphics.
        fig.savefig(output / f"{name}.svg", metadata={"Description": banner + "; " + font_note})
        fig.savefig(output / f"{name}.pdf", metadata={"Title": TITLES[name], "Subject": banner + "; " + font_note})
        plt.close(fig)


def _ppt_text(slide, text: str, x: float, y: float, w: float, h: float, *, size: int = 12, bold: bool = False, family: str = "DejaVu Serif"):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = line
        para.font.name = family
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = RGBColor.from_string("222222")
    return shape


def _disable_ppt_effects(presentation) -> None:
    """Suppress inherited theme shadows that PDF exporters rasterize."""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    for slide in presentation.slides:
        for shape in slide.shapes:
            for reference in shape._element.xpath(".//a:effectRef"):
                reference.set("idx", "0")
            properties = list(shape._element.xpath(".//p:spPr"))
            if shape.has_chart:
                properties.extend(shape.chart._chartSpace.xpath(".//c:spPr"))
            for prop in properties:
                for child in list(prop):
                    if child.tag in {qn("a:effectLst"), qn("a:effectDag")}:
                        prop.remove(child)
                prop.append(OxmlElement("a:effectLst"))


def _editable_ppt(aggregated: list[dict], output: Path, banner: str, font_note: str, typography: Typography) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(12), Inches(6.75)
    for figure in aggregated:
        name = figure["experiment"]
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _ppt_text(slide, TITLES[name], 0.35, 0.15, 11.3, 0.5, size=22, bold=True, family=typography.family)
        _ppt_text(slide, banner, 0.35, 0.7, 11.3, 0.42, size=11, bold=True, family=typography.family)
        _ppt_text(slide, _identity(figure), 0.35, 1.07, 11.3, 0.3, size=8, family=typography.family)
        _ppt_text(slide, _notes(name), 0.4, 5.8, 11.2, 0.45, size=10, family=typography.family)
        _ppt_text(slide, font_note, 0.4, 6.3, 11.2, 0.3, size=8, family=typography.family)
        if name == "A5":
            data = _comparison_rows(figure)
            shape = slide.shapes.add_table(len(data) + 1, 5, Inches(0.4), Inches(1.4), Inches(11.2), Inches(4.1))
            table = shape.table
            headers = ["Scheme / family", "Semantics / guarantee", "Security / parameters", "Evidence / source / scope", "Hardware / backend"]
            for r, row in enumerate([headers] + data):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = value
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string("DEDEDE" if r == 0 else ("F4F4F4" if r % 2 else "FFFFFF"))
                    for para in cell.text_frame.paragraphs:
                        para.font.name = typography.family
                        para.font.size = Pt(11 if len(data) <= 3 else 8)
                        para.font.bold = r == 0
                        para.font.color.rgb = RGBColor.from_string("222222")
            continue
        panels = _panels(figure)
        for panel_index, panel in enumerate(panels):
            x, width = (0.6, 10.9) if len(panels) == 1 else (0.4 + panel_index * 6, 5.3)
            y, height = 1.4, 4.1
            if not panel["series"]:
                # Native PowerPoint axes and text, with no dummy numeric series.
                for start_x, start_y, end_x, end_y in ((x + 0.45, y + 0.4, x + 0.45, y + 3.2), (x + 0.45, y + 3.2, x + width - 0.1, y + 3.2)):
                    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(start_x), Inches(start_y), Inches(end_x), Inches(end_y))
                    line.line.color.rgb = RGBColor.from_string("555555")
                _ppt_text(slide, panel["title"], x + 0.3, y - 0.15, width - 0.3, 0.4, size=12, bold=True, family=typography.family)
                _ppt_text(slide, "NO DATA\nAwaiting authorized local measurements", x + 0.8, y + 1.1, width - 1.1, 1.0, size=14, family=typography.family)
                _ppt_text(slide, panel["x_label"], x + 0.6, y + 3.25, width - 0.7, 0.5, size=11, family=typography.family)
                _ppt_text(slide, panel["y_label"], x + 0.55, y + 0.3, width - 0.7, 0.4, size=10, family=typography.family)
                failures = panel.get("failures", [])
                if failures:
                    _ppt_text(slide, "Tested failures: " + "; ".join(f"{f['status']} @ {f['x']:g} MiB" for f in failures), x + 0.6, y + 2.4, width - 0.7, 0.6, size=9, family=typography.family)
                continue
            if name == "A2":
                data = CategoryChartData()
                data.categories = [value.replace("\n", " / ") for value in panel["categories"]]
                for series in panel["series"]:
                    data.add_series(series["label"], series["y"])
                kind = XL_CHART_TYPE.COLUMN_STACKED
            else:
                data = XyChartData()
                for series in panel["series"]:
                    target = data.add_series(series["label"])
                    for x_val, y_val in zip(series["x"], series["y"]):
                        target.add_data_point(x_val, y_val)
                kind = XL_CHART_TYPE.XY_SCATTER_LINES
            chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(width), Inches(height), data).chart
            chart.has_title = True
            chart.chart_title.text_frame.text = panel["title"]
            chart.font.name = typography.family
            chart.font.size = Pt(10)
            chart.font.color.rgb = RGBColor.from_string("222222")
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(8)
            chart.category_axis.has_title = True
            chart.category_axis.axis_title.text_frame.text = panel["x_label"]
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = panel["y_label"]
            chart.value_axis.minimum_scale = 0
            for i, series in enumerate(chart.series):
                color = RGBColor.from_string(GRAY[i % len(GRAY)][1:])
                series.format.line.color.rgb = RGBColor.from_string("222222") if name == "A2" else color
                series.format.line.width = Pt(0.75 if name == "A2" else 1.25)
                if name == "A2":
                    series.format.fill.patterned()
                    series.format.fill.pattern = (MSO_PATTERN.PERCENT_90, MSO_PATTERN.DOWNWARD_DIAGONAL, MSO_PATTERN.UPWARD_DIAGONAL, MSO_PATTERN.DIAGONAL_CROSS, MSO_PATTERN.DOTTED_DIAMOND, MSO_PATTERN.CROSS, MSO_PATTERN.PERCENT_10)[i % 7]
                    series.format.fill.fore_color.rgb = RGBColor.from_string("444444")
                    series.format.fill.back_color.rgb = color
                else:
                    series.format.line.dash_style = (MSO_LINE_DASH_STYLE.SOLID, MSO_LINE_DASH_STYLE.DASH, MSO_LINE_DASH_STYLE.DASH_DOT, MSO_LINE_DASH_STYLE.ROUND_DOT)[i % 4]
                    series.marker.style = (XL_MARKER_STYLE.CIRCLE, XL_MARKER_STYLE.SQUARE, XL_MARKER_STYLE.TRIANGLE, XL_MARKER_STYLE.DIAMOND)[i % 4]
                    series.marker.size = 5
                    series.marker.format.fill.solid()
                    series.marker.format.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
                    series.marker.format.line.color.rgb = color
            failures = panel.get("failures", [])
            if failures:
                _ppt_text(slide, "Tested failures: " + "; ".join(f"{f['status']} @ {f['x']:g} MiB" for f in failures), x + 0.4, y + 0.4, width - 0.5, 0.4, size=9, family=typography.family)
            if panel.get("vram_missing"):
                _ppt_text(slide, "VRAM unavailable", x + 0.4, y + 0.8, width - 0.5, 0.3, size=9, family=typography.family)
    prs.core_properties.title = "Damascus A1–A5 layouts"
    prs.core_properties.subject = banner
    prs.core_properties.comments = "Native editable charts/tables/shapes. No screenshots. " + font_note
    _disable_ppt_effects(prs)
    prs.save(output / "A-group-editable.pptx")


def render(output: str | Path, records: list[dict] | None = None, *, fixture: bool = False, publication: bool = False) -> dict:
    """Render validated logs or empty layouts; fixture rendering is explicit opt-in."""
    records = [] if records is None else records
    if fixture and publication:
        raise ValidationError("fixture mode and publication mode are mutually exclusive")
    if any(r.get("data_kind") == "fixture" for r in records) and not fixture:
        raise ValidationError("fixture records require --fixture; publication use is forbidden")
    if fixture and any(r.get("data_kind") != "fixture" for r in records):
        raise ValidationError("fixture rendering cannot mix real or correctness records")
    aggregated = [aggregate(records, name, publication=publication) for name in EXPERIMENTS]
    typography = check_typography()
    if publication and not typography.publication_ready:
        raise ValidationError("publication typography unavailable: " + ", ".join(typography.missing))
    # Native PowerPoint text cannot embed TeX newtx math using python-pptx.
    # Keep all plot labels in plain text; no formula is passed off as newtx.
    banner, font_note = _status(aggregated, publication, typography)
    target = Path(output)
    target.mkdir(parents=True, exist_ok=True)
    for figure in aggregated:
        _vector(figure, target, banner, font_note, typography, publication)
    _editable_ppt(aggregated, target, banner, font_note, typography)
    manifest = {
        "schema_version": 1, "banner": banner, "publication": publication,
        "fixture": fixture, "typography": {"ready": typography.publication_ready, "family": typography.family, "missing": list(typography.missing)},
        "pptx": "Native editable charts, tables and axes. Plain-text formulas; no embedded screenshots; fonts are not embedded.",
        "vector": "SVG/PDF with vector paths/text; no rasterized plot images.",
        "limits": ["No data collection occurred.", "Artifact length/hash checks do not approve the unresolved protocol wire semantics.", "Epoch provenance/certification must be independently authenticated.", "PowerPoint uses editable plain text, not TeX newtx equation objects; final equation typography requires a separate approved workflow."],
        "figures": aggregated,
    }
    (target / "manifest.json").write_text(json.dumps(manifest, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--input", type=Path, help="JSONL/JSON evidence logs; omit for EMPTY layouts")
    parser.add_argument("--fixture", action="store_true", help="explicitly permit synthetic NON-PAPER fixture previews")
    parser.add_argument("--publication", action="store_true", help="reject fixtures and require exact publication typography")
    args = parser.parse_args()
    try:
        manifest = render(args.output, load_records(args.input) if args.input else [], fixture=args.fixture, publication=args.publication)
    except (ValidationError, OSError, subprocess.SubprocessError) as exc:
        parser.exit(2, f"plot export rejected: {exc}\n")
    print(json.dumps({"output": str(args.output), "banner": manifest["banner"], "typography": manifest["typography"]}, indent=2))


if __name__ == "__main__":
    main()
